# -*- coding: utf-8 -*-
"""
build_full_architecture_docx.py: 架构与实现设计说明书 (V1.2 含多租户与安全体系及架构优化蓝图) 主生成程序
"""

import os
import sys
import docx

# 添加当前目录至模块搜索路径
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from builder.common import (
    setup_document_page,
    add_title_block
)
from builder.part1_intro import build_part1
from builder.part2_business import build_part2
from builder.part3_finance_price import build_part3
from builder.part4_system_db import build_part4
from builder.part5_backup_license_tenant import build_part5
from builder.part6_optimization import build_part6

def main():
    target_docx = "睿链进销存与往来管理系统_架构与实现设计说明书_V1.2_含多租户与安全体系.docx"
    print(f"[*] 开始构建全面详尽的架构与实现设计说明书...")
    print(f"[*] 目标输出文件: {target_docx}")

    doc = docx.Document()
    setup_document_page(doc)

    # 1. 标题与文档控制元数据块
    add_title_block(
        doc,
        title="睿链进销存与往来管理系统",
        subtitle="系统架构、核心模块实现与全景架构演进优化设计说明书",
        version="V1.2（含多租户与安全体系及未来架构优化蓝图）",
        positioning="企业级采购、销售、库存及往来账务一体化管理平台",
        date_str="2026-10-15"
    )

    # 2. 第一至四章：引言、定位、技术架构与功能全景
    print(" -> 写入第一至四章（软件名称与品牌定位、项目定位与边界、总体技术架构、功能架构全景）...")
    build_part1(doc)

    # 3. 第五至八章：业务核心模块深度设计（基础资料、采购、销售、库存与仓储）
    print(" -> 写入第五至八章（基础资料、采购供应链、销售履约、库存仓储与成本引擎）...")
    build_part2(doc)

    # 4. 第九至十二章：财务资金与控制模块深度设计（往来账务、价格决策、预警中心、报表中心）
    print(" -> 写入第九至十二章（往来与资金核算、价格决策引擎、预警中心、报表与商业智能）...")
    build_part3(doc)

    # 5. 第十三至十五章：系统安全、数据库物理字典与单据状态机
    print(" -> 写入第十三至十五章（系统安全与RBAC、数据库表结构全景字典、单据状态机）...")
    build_part4(doc)

    # 6. 第十六至十八章：容灾备份、软件授权与多租户总控中枢
    print(" -> 写入第十六至十八章（备份与容灾、软件授权安全、SaaS多租户架构与总控中枢）...")
    build_part5(doc)

    # 7. 第十九至二十章：现有架构评估与未来架构优化蓝图
    print(" -> 写入第十九至二十章（现有架构瓶颈评估、面向未来的架构优化演进蓝图）...")
    build_part6(doc)

    # 保存文档
    doc.save(target_docx)
    print(f"[+] 文档构建成功并已保存至: {target_docx}")

    # 统计信息验证
    doc_verify = docx.Document(target_docx)
    p_count = len(doc_verify.paragraphs)
    t_count = len(doc_verify.tables)
    file_size_kb = os.path.getsize(target_docx) / 1024

    headings_count = 0
    words_est = 0
    for p in doc_verify.paragraphs:
        words_est += len(p.text)
        if any(p.text.startswith(prefix) for prefix in [
            "一、", "二、", "三、", "四、", "五、", "六、", "七、", "八、", "九、", "十",
            "1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.",
            "10.", "11.", "12.", "13.", "14.", "15.", "16.", "17.", "18.", "19.", "20."
        ]):
            headings_count += 1
    for t in doc_verify.tables:
        for r in t.rows:
            for c in r.cells:
                words_est += len(c.text)

    print(f"==================================================")
    print(f" 最终生成的 Word 文档统计分析报告:")
    print(f" - 文件大小: {file_size_kb:.1f} KB")
    print(f" - 段落总数: {p_count} 个段落")
    print(f" - 表格总数: {t_count} 个精心排版的专业表格")
    print(f" - 章节标题数: {headings_count} 个主要章节与子章节")
    print(f" - 核心中文字符数 (含表格): 约 {words_est} 字")
    print(f"==================================================")

if __name__ == "__main__":
    main()
