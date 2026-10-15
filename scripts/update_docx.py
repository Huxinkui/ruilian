# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def main():
    doc_path = '睿链进销存与往来管理系统_架构与实现设计说明书_V1.1_含授权体系.docx'
    doc = docx.Document(doc_path)

    # 1. 添加大标题
    h1 = doc.add_paragraph('二十四、SaaS 多租户架构演进、安全加固与平台总控中枢设计说明（V1.2 架构最新修改）', style='Heading 1')

    # 2. 24.1 架构演进背景与改造驱动
    doc.add_paragraph('24.1 架构演进背景与改造驱动', style='Heading 2')
    doc.add_paragraph(
        '随着系统从单机单租户向平台化、集团化多企业管控演进，系统需支撑多组织与多分支独立核算与业务自闭环。在公网和局域网部署环境下，传统企业管理系统常年遭受针对通用用户名（admin/administrator）的大规模自动化爬虫扫描、弱密码探测与字典暴力破解攻击。为确保平台达到高等级安全防护标准，必须消除“默认管理员暴露面”，并实现租户数据物理级隔离与管理权限的强边界约束。'
    )
    doc.add_paragraph('本次架构演进聚焦三大核心诉求：')
    p = doc.add_paragraph('SaaS 多租户物理级强隔离：彻底摒弃逻辑隔离，消除 SQL 拼装漏加租户标识所带来的数据穿透隐患；', style='List Bullet')
    p = doc.add_paragraph('平台管理员专属化与攻击面消除：注销默认 admin 账号，启用专属标识 huxinkui，配备 IP 暴力破解防御与图形验证码；', style='List Bullet')
    p = doc.add_paragraph('管理权界强隔离与差异化交互：普通租户全面隐藏多租户管理页面并施加路由与接口强拦截；为超级管理员重新设计全生命周期的租户总控平台控制台。', style='List Bullet')

    # 3. 24.2 Database-per-Tenant 物理数据库独立架构
    doc.add_paragraph('24.2 Database-per-Tenant 物理数据库独立架构', style='Heading 2')
    doc.add_paragraph(
        '为确保各企业租户数据的绝对隔离与安全独立，系统采用 Database-per-Tenant 物理数据库独立架构：'
    )
    doc.add_paragraph('物理沙箱设计：平台主管理库（data/ruilian.db）负责存储全局多租户注册信息（tenants 表）、全局系统配置、超级管理员专属账号、主空间审计与全量备份元数据；子租户独立库（data/tenants/{tenant_code}/ruilian.db）则为每个租户分配专属物理 SQLite 数据库。连接池、WAL 读写日志、事务锁与 Schema 完全物理隔离，彻底杜绝数据越权与数据泄露隐患。')
    doc.add_paragraph('动态租户路由中间件：通过 HTTP 请求头 X-Tenant-Code、Cookie ruilian_tenant 或 Query 参数动态分发请求。TenantManager 内部维护租户 App 实例缓存池，按需动态加载并保证线程并发安全。')
    doc.add_paragraph('物理存储配额与写操作拦截引擎（Storage Quota Engine）：各租户拥有独立存储配额（max_storage_mb，支持 200MB、500MB、1GB、2GB 或不限容量）。系统精准穿透统计主库文件、WAL 文件（-wal）、共享内存文件（-shm）及日志目录总物理体积。在所有写操作（POST/PUT/DELETE）进入业务层前执行 CheckStorageQuota 校验，超限即刻返回 HTTP 403 (STORAGE_QUOTA_EXCEEDED)，只保留读权限，防止恶意膨胀击穿服务器磁盘。')

    # 4. 24.3 核心安全防线与管理员账号专属化
    doc.add_paragraph('24.3 核心安全防线与管理员账号专属化（admin -> huxinkui）', style='Heading 2')
    doc.add_paragraph(
        '废弃默认 admin 账号防范定向攻击：彻底注销与拔除系统各层级默认 admin 账号的合法身份。任何使用 admin 的登录尝试均直接拒绝报错，抹平众所周知的公网字典扫描面。将系统唯一超级管理员账号升级为专属标识 huxinkui，作为系统后台最高权限管控凭证。'
    )
    doc.add_paragraph('平滑自动化迁移机制：系统启动阶段，在 EnsureDefaultAccount 与 ensureTenantsTable 中内嵌自动迁移指令，将主数据库中现存的 admin 用户无损升级为 huxinkui，密码哈希、会话及权限完全平移。')
    doc.add_paragraph('租户空间身份互斥隔离：各子租户数据库中严格禁止存在 admin 或 huxinkui 账号，初始化时执行强力清退。各租户拥有与其租户代码相同的专属管理员账号（如 xuyonghong），租户无法以系统超级管理员身份登录自身空间。')
    doc.add_paragraph('IP 防爆破拦截器与图形验证码防御体系：IPLoginLimiter 针对来源 IP 进行登录频次与失败统计，单 IP 连续失败 8 次自动触发 15 分钟临时封禁；CaptchaStore 内存化安全验证码结合字符扰动、干扰线与动态盐值，阻断自动化脚本攻击。')

    # 5. 24.4 权限边界与多租户管理中枢隔离
    doc.add_paragraph('24.4 权限边界与多租户管理中枢隔离', style='Heading 2')
    doc.add_paragraph('前端租户界面完全屏蔽：普通租户登录后，侧边栏导航中绝不渲染「多租户管理」菜单项；隐藏多租户切换选择器，租户仅感知自身企业名称与空间使用量，无感知其他租户存在。')
    doc.add_paragraph('前端路由全天候强守卫（Navigation Guard）：Vue Router 全局前置守卫拦截对 /tenants 的所有访问，非超级管理员（huxinkui）访问立即阻断、弹出警告通知并重定向回业务首页。')
    doc.add_paragraph('后端 API 强鉴权中间件（requireMasterAdmin）：/api/system/tenants 旗下所有接口全量挂载 requireMasterAdmin。强制校验会话用户必须满足 u.IsMasterAdmin || u.Username == "huxinkui" 且当前处于主平台空间，任何子租户发起跨空间请求一律返回 403 拒绝，彻底阻断水平越权与垂直越权。')

    # 6. 24.5 超级管理员专属平台控制台（Admin Dashboard）
    doc.add_paragraph('24.5 超级管理员专属平台控制台（Admin Dashboard）设计', style='Heading 2')
    doc.add_paragraph('页面架构双模态重构：普通租户进入标准进销存业务经营看板（经营指标、采购走势、库存结构、应收账龄、常用单据开单快捷方式）；超级管理员则直接呈现全新设计的平台总控与多租户管理运维中心（AdminDashboard.vue）。')
    doc.add_paragraph('总控中枢核心界面板块：')
    doc.add_paragraph('1. 平台运维 Hero 中枢：超级管理员专属标识、系统健康度实时脉冲指示灯、多租户全貌调度；', style='List Bullet')
    doc.add_paragraph('2. 四维运维 KPI 网格：租户空间总数及健康状态、全局已用物理存储及配额使用率、容量预警租户数（≥80% 需扩容提示）、核心防御运行状态；', style='List Bullet')
    doc.add_paragraph('3. 租户空间管控舰队（Tenant Fleet）：支持快速搜索、状态过滤与配额进度条（支持危险预警色），透视物理存储占用与路径（data/tenants/{code}/ruilian.db），提供一键跨空间巡检、动态配额调整、租户密码重置、启停状态切换等快捷操作；', style='List Bullet')
    doc.add_paragraph('4. ECharts 数据可视化看板：各租户物理数据库体积占用比例与配额消耗图表；', style='List Bullet')
    doc.add_paragraph('5. 跨空间监管巡检体验：超级管理员在管理后台可随时“一键进入”任意租户空间进行巡检审计；进入后显示红底监管横幅并附带“返回平台总控台”按钮，实现跨空间监管与平台中枢的顺畅穿透。', style='List Bullet')

    # 7. 24.6 架构变更版本总结
    doc.add_paragraph('24.6 架构版本总结', style='Heading 2')
    doc.add_paragraph(
        '版本说明：V1.2 在 V1.1 架构基础上，全面实现 SaaS Database-per-Tenant 物理数据库独立架构演进、平台超级管理员专属账号改造（admin -> huxinkui）及防暴力破解安全加固、租户管理页面与接口强隔离，以及超级管理员专属运维控制台（Admin Dashboard）的全新设计。'
    )

    # 保存原文档和新命名文档
    doc.save(doc_path)
    new_doc_path = '睿链进销存与往来管理系统_架构与实现设计说明书_V1.2_含多租户与安全体系.docx'
    doc.save(new_doc_path)
    print(f'Successfully updated: {doc_path} and created: {new_doc_path}')

if __name__ == '__main__':
    main()
