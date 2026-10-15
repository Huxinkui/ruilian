# -*- coding: utf-8 -*-
"""
common.py: 统一文档样式、排版工具函数与表格渲染组件
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

FONT_NAME_CN = "Microsoft YaHei"
FONT_NAME_EN = "Segoe UI"
FONT_NAME_CODE = "Consolas"

# 调色板定义
COLOR_PRIMARY = RGBColor(30, 58, 138)      # 深海蓝 #1E3A8A (H1)
COLOR_SECONDARY = RGBColor(3, 105, 161)    # 蔚蓝 #0369A1 (H2)
COLOR_TERTIARY = RGBColor(51, 65, 85)      # 板岩暗灰 #334155 (H3)
COLOR_TEXT = RGBColor(31, 41, 55)          # 正文灰黑 #1F2937
COLOR_MUTED = RGBColor(100, 116, 139)      # 辅助灰 #64748B
COLOR_ACCENT = RGBColor(220, 38, 38)       # 警告红 #DC2626

def set_font_run(run, font_name=FONT_NAME_CN, size_pt=10.5, color=COLOR_TEXT, bold=False, italic=False):
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:eastAsia="{font_name}"/>')
    rPr.append(rFonts)
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic

def setup_document_page(doc):
    for sec in doc.sections:
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11.0)
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

def add_title_block(doc, title, subtitle, version, positioning, date_str):
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(18)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run(title)
    set_font_run(r_title, font_name=FONT_NAME_CN, size_pt=22, color=COLOR_PRIMARY, bold=True)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run(subtitle)
    set_font_run(r_sub, font_name=FONT_NAME_CN, size_pt=13, color=COLOR_SECONDARY, bold=True)

    # 元数据信息卡片
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta = [
        [f"文档版本：{version}", f"发布日期：{date_str}"],
        [f"系统定位：{positioning}", "安全等级：企业级物理隔离与防篡改体系"]
    ]
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = meta[r_idx][c_idx]
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
            cell._tc.get_or_add_tcPr().append(shd)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                set_font_run(r, font_name=FONT_NAME_CN, size_pt=9.5, color=COLOR_MUTED, bold=(r_idx==0 and c_idx==0))
    p_sep = doc.add_paragraph()
    p_sep.paragraph_format.space_after = Pt(16)

def add_h1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    set_font_run(run, font_name=FONT_NAME_CN, size_pt=15, color=COLOR_PRIMARY, bold=True)
    return h

def add_h2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    set_font_run(run, font_name=FONT_NAME_CN, size_pt=12.5, color=COLOR_SECONDARY, bold=True)
    return h

def add_h3(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(3)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    set_font_run(run, font_name=FONT_NAME_CN, size_pt=11, color=COLOR_TERTIARY, bold=True)
    return h

def add_p(doc, text, bold_prefix="", space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        set_font_run(r_pre, font_name=FONT_NAME_CN, size_pt=10.5, color=COLOR_TEXT, bold=True)
    r = p.add_run(text)
    set_font_run(r, font_name=FONT_NAME_CN, size_pt=10.5, color=COLOR_TEXT)
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        set_font_run(r_pre, font_name=FONT_NAME_CN, size_pt=10, color=COLOR_TEXT, bold=True)
    r = p.add_run(text)
    set_font_run(r, font_name=FONT_NAME_CN, size_pt=10, color=COLOR_TEXT)
    return p

def add_callout(doc, text, title="设计原则与关键规则", bg_color="EFF6FF", border_color="2563EB"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
    tcPr.append(shd)
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="120" w:type="dxa"/>
            <w:bottom w:w="120" w:type="dxa"/>
            <w:left w:w="180" w:type="dxa"/>
            <w:right w:w="180" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r_t = p.add_run(f"【{title}】\n")
    set_font_run(r_t, font_name=FONT_NAME_CN, size_pt=10, color=COLOR_PRIMARY, bold=True)
    r_b = p.add_run(text)
    set_font_run(r_b, font_name=FONT_NAME_CN, size_pt=9.5, color=COLOR_TEXT)

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
    tcPr.append(shd)
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="94A3B8"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="100" w:type="dxa"/>
            <w:bottom w:w="100" w:type="dxa"/>
            <w:left w:w="150" w:type="dxa"/>
            <w:right w:w="150" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(code_text.strip())
    set_font_run(r, font_name=FONT_NAME_CODE, size_pt=8.5, color=RGBColor(30, 41, 59))

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_after = Pt(4)

def add_table(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 表头格式化
    for c_idx, h_text in enumerate(headers):
        cell = table.rows[0].cells[c_idx]
        cell.text = h_text
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E3A8A"/>')
        tcPr.append(shd)
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="100" w:type="dxa"/><w:bottom w:w="100" w:type="dxa"/><w:left w:w="120" w:type="dxa"/><w:right w:w="120" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            set_font_run(r, font_name=FONT_NAME_CN, size_pt=9.5, color=RGBColor(255, 255, 255), bold=True)

    # 数据行格式化
    for r_idx, row in enumerate(rows_data):
        is_even = (r_idx % 2 == 1)
        bg = "F8FAFC" if is_even else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            tcPr = cell._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}"/>')
            tcPr.append(shd)
            tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="80" w:type="dxa"/><w:bottom w:w="80" w:type="dxa"/><w:left w:w="100" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tcMar>')
            tcPr.append(tcMar)
            p = cell.paragraphs[0]
            for r in p.runs:
                set_font_run(r, font_name=FONT_NAME_CN, size_pt=9.0, color=COLOR_TEXT)

    # 细边框设置
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

    # 宽度设置
    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_after = Pt(6)
